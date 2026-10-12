"""Write the ExamGrid datasets into one formatted Excel workbook."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
ASSUMED_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
# Identifier columns stay text even when they look numeric (e.g. roll code "05").
TEXT_COLUMNS = {"roll_no", "roll_code", "subject_code", "room_id", "faculty_id", "section"}

# workbook sheet -> (table key, description)
SHEETS = [
    ("Students", "students", "One row per student: roll number, branch, year, section, lateral entry, ground-floor need."),
    ("Subjects", "subjects", "SRKR R23 written-exam courses for this odd semester, incl. PE and open electives."),
    ("Registrations", "registrations", "Every paper each student writes this cycle (REGULAR or BACKLOG)."),
    ("Rooms", "rooms", "Exam rooms with seat grid (rows x cols), floor and availability."),
    ("Faculty", "faculty", "Faculty with department, designation and maximum invigilation duties."),
    ("Faculty_Subjects", "faculty_subjects", "Who teaches which subject (they must not invigilate it)."),
    ("Faculty_Leave", "faculty_unavailability", "Faculty leave during the exam window (FN / AN / BOTH)."),
    ("Attendance_History", "attendance_history", "Past turnout, Nov 2022-2025: training data for the turnout model."),
    ("Branches", "branches", "Branch codes and 2026 intake. Yellow = roll code assumed, please confirm."),
]


def _num(v):
    """CSV-style strings that are numbers become numbers, so Excel can sum them."""
    if isinstance(v, str):
        if len(v) > 1 and v.startswith("0") and not v.startswith("0."):
            return v  # codes like "05" must keep their leading zero
        try:
            return int(v) if v.lstrip("-").isdigit() else float(v) if v.replace(".", "", 1).isdigit() else v
        except ValueError:
            return v
    return v


def _table(ws, rows: list[dict], headers: list[str] | None = None) -> None:
    headers = headers or (list(rows[0]) if rows else [])
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) if h in TEXT_COLUMNS else _num(r.get(h)) for h in headers])
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body = Font(name=FONT, size=10)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body
    ws.freeze_panes = "A2"
    if headers:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    sample = rows[:300]
    for i, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(r.get(h, ""))) for r in sample])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(9, longest + 2))
    ws.row_dimensions[1].height = 30


def _about(ws, config: dict, scale: str, seed: int, counts: dict) -> None:
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    lines = [
        ("ExamGrid - SRKR Engineering College datasets", None),
        (config["college"], None),
        ("", None),
        ("Exam cycle", config["exam_cycle"]),
        ("Regulation", f"{config['regulation']} (with effect from 2023-24 admitted batch)"),
        ("Exam window", f"{config['start_date']} to {config['last_allowed_date']}, no Sundays, holiday(s): "
                        + ", ".join(config["holidays"])),
        ("Sessions", "; ".join(f"{s['code']} {s['start']}-{s['end']}" for s in config["sessions"])),
        ("Dataset scale / seed", f"{scale} / {seed}  (regenerate: python scripts/generate_datasets.py --scale {scale})"),
        ("", None),
        ("SHEETS", None),
    ]
    for sheet, key, desc in SHEETS:
        lines.append((sheet, f"{desc}  [{counts[key]:,} rows]"))
    lines += [
        ("Summary", "Counts per branch and year, calculated with formulas from the sheets above."),
        ("", None),
        ("WHAT IS REAL", None),
        ("Curriculum", "Course codes, names, categories, credits, CIE/SEE marks, PE options and open electives "
                       "are taken from SRKR's official R23 syllabus PDFs (https://srkrec.ac.in/syllabus/)."),
        ("Roll numbers", "Format 23B91A0512 (regular) / 24B95A0503 (lateral entry), from SRKR's published result lists."),
        ("Intake", "2026 sanctioned intake per branch (Collegedunia)."),
        ("", None),
        ("WHAT IS SYNTHETIC", None),
        ("People", "Every student and faculty name, roll-number assignment and ID is invented. No real person's data."),
        ("Rooms & attendance", "Rooms, blocks, leave entries and attendance history are invented but realistic."),
        ("Assumptions", "Roll codes for AIML (61), CSIT (59), CSD (44), CIC (49) are assumed - highlighted "
                        "yellow on the Branches sheet. Correct them in scripts/generate_datasets.py (BRANCHES)."),
        ("", None),
        ("HOW TO READ", None),
        ("Header row", "Dark blue. Every data sheet has filters and a frozen header row."),
        ("Yellow cells", "Assumed values that should be confirmed with the SRKR exam cell."),
    ]
    for i, (a, b) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=a).font = Font(name=FONT, size=10, bold=b is None or i <= 2)
        if b is not None:
            c = ws.cell(row=i, column=2, value=b)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="1F3864")
    ws["A2"].font = Font(name=FONT, size=11, italic=True)


def _summary(ws, branches: list[dict], n_students: int, n_regs: int, n_rooms: int, n_faculty: int) -> None:
    """Live formulas over the data sheets, so the numbers update if rows are edited."""
    st, rg, rm, fc = n_students + 1, n_regs + 1, n_rooms + 1, n_faculty + 1
    headers = ["Branch", "Year 1", "Year 2", "Year 3", "Year 4", "Total students", "Lateral entry",
               "Need ground floor", "2026 intake", "Faculty"]
    ws.append(headers)
    for i, b in enumerate(branches, start=2):
        name = b["branch"]
        row = [name]
        for col, year in zip("BCDE", range(1, 5)):
            row.append(f'=COUNTIFS(Students!$D$2:$D${st},$A{i},Students!$E$2:$E${st},{year})')
        row += [f"=SUM(B{i}:E{i})",
                f'=COUNTIFS(Students!$D$2:$D${st},$A{i},Students!$H$2:$H${st},"LATERAL")',
                f'=COUNTIFS(Students!$D$2:$D${st},$A{i},Students!$J$2:$J${st},"yes")',
                f"=INDEX(Branches!$E:$E,MATCH($A{i},Branches!$A:$A,0))",
                f'=COUNTIF(Faculty!$C$2:$C${fc},$A{i})']
        ws.append(row)
    last = len(branches) + 1
    total = last + 1
    ws.append(["TOTAL"] + [f"=SUM({c}2:{c}{last})" for c in "BCDEFGHIJ"])
    ws.cell(row=total, column=10, value=f"=SUM(J2:J{last})+COUNTIF(Faculty!$C$2:$C${fc},\"BSH\")")
    for cell in ws[total]:
        cell.font = Font(name=FONT, bold=True)
        cell.fill = TOTAL_FILL

    r = total + 2
    facts = [
        ("Registrations (papers to sit)", f"=COUNTA(Registrations!$A$2:$A${rg})"),
        ("  of which REGULAR", f'=COUNTIF(Registrations!$C$2:$C${rg},"REGULAR")'),
        ("  of which BACKLOG", f'=COUNTIF(Registrations!$C$2:$C${rg},"BACKLOG")'),
        ("Backlog share", f"=IFERROR(B{r + 2}/B{r},0)"),
        ("Exam-usable rooms", f'=COUNTIF(Rooms!$I$2:$I${rm},"yes")'),
        ("Exam-usable seats", f'=SUMIF(Rooms!$I$2:$I${rm},"yes",Rooms!$F$2:$F${rm})'),
        ("Ground-floor usable seats", f'=SUMIFS(Rooms!$F$2:$F${rm},Rooms!$I$2:$I${rm},"yes",Rooms!$H$2:$H${rm},"yes")'),
        ("Faculty in Basic Sciences & Humanities", f'=COUNTIF(Faculty!$C$2:$C${fc},"BSH")'),
    ]
    for j, (label, formula) in enumerate(facts):
        ws.cell(row=r + j, column=1, value=label).font = Font(name=FONT, size=10, bold=not label.startswith("  "))
        c = ws.cell(row=r + j, column=2, value=formula)
        c.font = Font(name=FONT, size=10)
        if label == "Backlog share":
            c.number_format = "0.0%"
        else:
            c.number_format = "#,##0"

    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.number_format = "#,##0"
            cell.border = Border(bottom=THIN)
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = "B2"


def export_workbook(path: Path, tables: dict, config: dict, scale: str, seed: int) -> Path:
    wb = Workbook()
    about = wb.active
    about.title = "About"
    counts = {k: len(v) for k, v in tables.items()}
    _about(about, config, scale, seed, counts)

    _summary(wb.create_sheet("Summary"), tables["branches"], counts["students"], counts["registrations"],
             counts["rooms"], counts["faculty"])

    for sheet, key, _ in SHEETS:
        ws = wb.create_sheet(sheet)
        headers = ["faculty_id", "date", "session", "reason"] if key == "faculty_unavailability" and not tables[key] else None
        _table(ws, tables[key], headers)
        if key == "branches":
            status_col = list(tables[key][0]).index("roll_code_status") + 1
            for row in ws.iter_rows(min_row=2):
                if str(row[status_col - 1].value).startswith("ASSUMED"):
                    for cell in row[1:status_col]:
                        cell.fill = ASSUMED_FILL

    cfg = wb.create_sheet("Exam_Config")
    cfg.append(["setting", "value"])
    for k, v in config.items():
        if isinstance(v, list):
            v = "; ".join(f"{s['code']} {s['start']}-{s['end']}" if isinstance(s, dict) else str(s) for s in v)
        cfg.append([k, v])
    _table_style = cfg[1]
    for cell in _table_style:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in cfg.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
    cfg.column_dimensions["A"].width = 34
    cfg.column_dimensions["B"].width = 70

    wb.calculation.fullCalcOnLoad = True  # Excel computes the Summary formulas on open
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
